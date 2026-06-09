from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Nivaana_Ecommerce_Image_Aspect_Ratio_Spec.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips=9360):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_twips))


def style_cell(cell, header=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "F3F4F6")
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9.5 if not header else 9)
            if header:
                run.font.bold = True


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    hdr_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        hdr_cells[idx].text = header
        style_cell(hdr_cells[idx], header=True)
        hdr_cells[idx].width = Inches(widths[idx])

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            style_cell(cells[idx])
            cells[idx].width = Inches(widths[idx])
    return table


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].font.color.rgb = RGBColor(31, 41, 55)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 2"].font.bold = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "Nivaana E-Commerce Image Ratio Spec"
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.runs[0].font.name = "Arial"
    header_p.runs[0].font.size = Pt(8)
    header_p.runs[0].font.color.rgb = RGBColor(107, 114, 128)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.text = "Prepared for Nivaana Web"
    footer_p.runs[0].font.name = "Arial"
    footer_p.runs[0].font.size = Pt(8)
    footer_p.runs[0].font.color.rgb = RGBColor(107, 114, 128)

    title = doc.add_paragraph(style="Title")
    title.add_run("Nivaana Image Aspect Ratio Specification")
    subtitle = doc.add_paragraph()
    subtitle.add_run(
        "Recommended upload/export ratios for product cards, Deal of the Day carousel, category cards, product detail, and hero media."
    ).italic = True

    doc.add_heading("Executive Recommendation", level=1)
    p = doc.add_paragraph()
    p.add_run("Use 5:4 as the master product image ratio. ").bold = True
    p.add_run(
        "Export product images at 1600 x 1280 px where possible, with the product kept inside the center 80% safe area. This matches the current product card crop and also works in Deal of the Day because that carousel displays images with object-contain."
    )

    add_bullet(doc, "Products: 5:4, recommended 1600 x 1280 px.")
    add_bullet(doc, "Category/fragrance thumbnails: 1:1, recommended 1200 x 1200 px, unless the UI is changed to reuse 5:4.")
    add_bullet(doc, "Hero desktop: wide 21:9-ish media, recommended 2560 x 1160 px or 2400 x 1080 px.")
    add_bullet(doc, "Hero mobile: separate 1:1 or 4:5 asset, recommended 1200 x 1200 px or 1080 x 1350 px.")

    doc.add_heading("Current App Behavior", level=1)
    add_table(
        doc,
        ["Placement", "Current container / fit", "What this means", "Best asset ratio"],
        [
            [
                "Product cards",
                "aspect-[4/3.2], object-cover",
                "The card crops to roughly 5:4. Images wider/taller than this can look zoomed.",
                "5:4",
            ],
            [
                "Compact product cards / new arrivals",
                "aspect-[4/3.1], object-cover",
                "Very close to 5:4. A 5:4 master image will fit cleanly.",
                "5:4",
            ],
            [
                "Shop by fragrance/category cards",
                "aspect-square, object-cover",
                "This is why reused product images can crop heavily in category tiles.",
                "1:1, or change UI to 5:4/object-contain",
            ],
            [
                "Deal of the Day carousel",
                "variable wide card, object-contain",
                "The full product image is preserved, so 5:4 works without cropping.",
                "5:4 product master",
            ],
            [
                "Product detail main image",
                "fixed height, object-contain",
                "Flexible. It preserves the full image and does not force crop.",
                "5:4 or 1:1",
            ],
            [
                "Hero media",
                "large responsive frame, object-cover/object-contain by slide",
                "Desktop and mobile need different safe areas because the frame changes shape.",
                "Desktop 21:9-ish, mobile 1:1 or 4:5",
            ],
        ],
        [1.45, 1.8, 2.75, 1.2],
    )

    doc.add_heading("Recommended Upload Specs", level=1)
    add_table(
        doc,
        ["Asset type", "Ratio", "Recommended pixels", "Safe-area guidance"],
        [
            ["Product master image", "5:4", "1600 x 1280 px", "Keep product inside center 80%; leave background padding."],
            ["Minimum product image", "5:4", "1200 x 960 px", "Acceptable for catalog cards and detail pages."],
            ["Category thumbnail", "1:1", "1200 x 1200 px", "Compose specifically for square crop; avoid tall product close-ups."],
            ["Deal carousel", "5:4", "1600 x 1280 px", "Reuse product master; carousel uses contain so full pack is visible."],
            ["Hero desktop", "~21:9", "2560 x 1160 px", "Keep important subject center-left/center; avoid baked-in text."],
            ["Hero mobile", "1:1 or 4:5", "1200 x 1200 or 1080 x 1350 px", "Subject should remain readable in a narrow viewport."],
        ],
        [1.55, 1.0, 1.45, 3.2],
    )

    doc.add_heading("Why Some Images Look Too Zoomed", level=1)
    p = doc.add_paragraph()
    p.add_run("The zoom is not random. ").bold = True
    p.add_run(
        "The product cards and category cards use object-cover, which fills the container by cropping any extra width or height. The Deal of the Day carousel and product detail image use object-contain, which preserves the entire image but can leave empty space around it. If one image is reused everywhere, it will naturally behave differently across these sections."
    )

    doc.add_heading("Practical Production Rule", level=1)
    add_bullet(doc, "Shoot/export every product in 5:4 first. This becomes the catalog master.")
    add_bullet(doc, "Do not crop product packaging tightly. Leave air around the product on all sides.")
    add_bullet(doc, "Create separate 1:1 category thumbnails only for category/fragrance carousel tiles.")
    add_bullet(doc, "Create separate desktop and mobile hero images/videos. Do not rely on one hero asset for every screen.")
    add_bullet(doc, "Avoid text inside images for product and hero assets; keep text in the website layer.")

    doc.add_heading("Optional UI Adjustment", level=1)
    doc.add_paragraph(
        "If the goal is to reuse the same product image everywhere, change the Shop by fragrance/category card image area from square object-cover to either a 5:4 product-card ratio or object-contain. That would reduce the heavy crop shown in the Kasturi category card screenshot."
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
